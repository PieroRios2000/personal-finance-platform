"""`forecasting.plan_import.check_plan`: what makes a plan workbook importable, and
that every problem is reported at once with a row number or a field name, never a
value (T57, ADR 0048). All data is synthetic."""

from datetime import date, datetime
from pathlib import Path

import pytest
from forecasting.plan_import import Goal, check_plan

from forecasting.plan_file import PlanFile, PlanItem, read_plan, write_plan

TODAY = date(2026, 10, 4)
ACCOUNTS = {"BCP", "Ripley"}


def _item(row: int = 2, **overrides: object) -> PlanItem:
    fields: dict[str, object] = {
        "bank": "BCP",
        "description": f"PLANTED ITEM {row}",
        "currency": "PEN",
        "category": "Servicios",
        "months_seen": 6,
        "typical_amount": 50.0,
        "proposed_kind": "fixed",
        "kind": "fixed",
        "expected_amount": 50.0,
        "row": row,
    }
    fields.update(overrides)
    return PlanItem(**fields)  # type: ignore[arg-type]


def _meta(**overrides: object) -> dict[str, object]:
    meta: dict[str, object] = {
        "goal_amount": 10000,
        "usd_to_pen": 3.75,
        "emergency_months": 6,
        "emergency_basis": "all",
        "emergency_account": "Ripley",
        "target_date": datetime(2027, 12, 31),
        "income_pen_override": None,
        "income_usd_override": None,
    }
    meta.update(overrides)
    return meta


def _plan(items: list[PlanItem] | None = None, **meta: object) -> PlanFile:
    return PlanFile(tuple(items if items is not None else [_item()]), _meta(**meta))


def _check(plan: PlanFile) -> tuple[Goal | None, list[str]]:
    return check_plan(plan, today=TODAY, asset_accounts=ACCOUNTS)


def test_a_good_plan_has_no_problems_and_yields_the_goal() -> None:
    goal, problems = _check(_plan())

    assert problems == []
    assert goal == Goal(
        goal_amount=10000.0,
        usd_to_pen=3.75,
        emergency_months=6,
        emergency_basis="all",
        emergency_account="Ripley",
        target_date=date(2027, 12, 31),
        income_pen_override=None,
        income_usd_override=None,
    )


def test_optional_fields_may_be_empty_and_a_text_date_is_read() -> None:
    goal, problems = _check(_plan(target_date="2027-06-30", income_pen_override=4200))

    assert problems == []
    assert goal is not None
    assert goal.target_date == date(2027, 6, 30)
    assert goal.income_pen_override == 4200.0
    goal, problems = _check(_plan(target_date=None))
    assert problems == [] and goal is not None and goal.target_date is None


@pytest.mark.parametrize("kind", ["fixed", "variable", "ignore"])
def test_every_known_kind_is_accepted(kind: str) -> None:
    _, problems = _check(_plan([_item(kind=kind)]))

    assert problems == []


def test_an_unknown_or_empty_kind_names_the_row() -> None:
    _, problems = _check(_plan([_item(2), _item(3, kind="fijo"), _item(4, kind="")]))

    assert problems == [
        "Gastos fijos: kind in row 3 must be fixed, variable or ignore",
        "Gastos fijos: kind in row 4 must be fixed, variable or ignore",
    ]


def test_a_fixed_item_needs_a_positive_expected_amount() -> None:
    items = [
        _item(2, expected_amount=None),
        _item(3, expected_amount=0.0),
        _item(4, expected_amount=-5.0),
        _item(5, expected_amount=float("nan")),
        _item(6, expected_amount=float("inf")),
    ]

    _, problems = _check(_plan(items))

    assert problems == [
        f"Gastos fijos: expected_amount in row {n} must be a positive number"
        for n in range(2, 7)
    ]


def test_a_variable_item_may_leave_the_amount_empty_but_not_negative() -> None:
    _, ok = _check(_plan([_item(2, kind="variable", expected_amount=None)]))
    _, bad = _check(_plan([_item(2, kind="ignore", expected_amount=-1.0)]))

    assert ok == []
    assert bad == ["Gastos fijos: expected_amount in row 2 must be a positive number"]


def test_an_unknown_currency_and_a_missing_bank_name_the_row() -> None:
    _, problems = _check(_plan([_item(2, currency="EUR"), _item(3, bank="")]))

    assert problems == [
        "Gastos fijos: currency in row 2 must be PEN or USD",
        "Gastos fijos: bank in row 3 is empty",
    ]


def test_the_same_item_twice_is_reported_on_the_second_row() -> None:
    twin = _item(3, description=_item(2).description)

    _, problems = _check(_plan([_item(2), twin]))

    assert problems == [
        "Gastos fijos: row 3 repeats the bank, description and currency of row 2"
    ]


@pytest.mark.parametrize("field", ["goal_amount", "usd_to_pen", "emergency_months"])
def test_required_meta_fields_cannot_be_empty(field: str) -> None:
    _, problems = _check(_plan(**{field: None}))

    assert problems == [f"Meta: {field} is required"]


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("goal_amount", 0),
        ("goal_amount", -100),
        ("goal_amount", "diez mil"),
        ("usd_to_pen", 0),
        ("usd_to_pen", "3,75"),
        ("income_pen_override", -1),
        ("income_usd_override", "mucho"),
    ],
)
def test_meta_numbers_must_be_positive_numbers(field: str, value: object) -> None:
    _, problems = _check(_plan(**{field: value}))

    assert problems == [f"Meta: {field} must be a positive number"]


@pytest.mark.parametrize("months", [0, 25, 6.5, "seis"])
def test_emergency_months_is_a_whole_number_from_1_to_24(months: object) -> None:
    _, problems = _check(_plan(emergency_months=months))

    assert problems == ["Meta: emergency_months must be a whole number from 1 to 24"]


def test_emergency_months_accepts_the_edges_and_a_float_that_is_whole() -> None:
    for months in (1, 24, 6.0):
        goal, problems = _check(_plan(emergency_months=months))
        assert problems == [] and goal is not None
        assert goal.emergency_months == int(months)


def test_emergency_basis_must_be_one_of_two_values() -> None:
    _, problems = _check(_plan(emergency_basis="most"))

    assert problems == ["Meta: emergency_basis must be all or fixed_only"]


def test_the_target_date_must_be_a_date_after_today() -> None:
    _, past = _check(_plan(target_date=datetime(2026, 10, 4)))
    _, garbage = _check(_plan(target_date="pronto"))

    assert past == ["Meta: target_date must be a date after today"]
    assert garbage == ["Meta: target_date must be a date after today"]


def test_the_emergency_account_must_be_an_asset_account_the_lake_knows() -> None:
    _, problems = _check(_plan(emergency_account="ripley"))
    _, empty = _check(_plan(emergency_account=None))

    assert problems == [
        "Meta: emergency_account does not match an asset account "
        "(known accounts: BCP, Ripley)"
    ]
    assert empty == ["Meta: emergency_account is required"]


def test_all_problems_come_at_once_and_no_goal_is_returned() -> None:
    plan = _plan(
        [_item(2, kind="x"), _item(3, expected_amount=None)],
        goal_amount=-1,
        usd_to_pen=None,
    )

    goal, problems = _check(plan)

    assert goal is None
    assert len(problems) == 4


def test_no_problem_message_ever_contains_a_value_the_owner_typed() -> None:
    plan = _plan(
        [_item(2, description="SECRET SHOP", kind="PRIVATE-KIND", currency="ZZZ")],
        goal_amount="PRIVATE-AMOUNT",
        emergency_account="PRIVATE-ACCOUNT",
        target_date="PRIVATE-DATE",
    )

    _, problems = _check(plan)

    assert problems
    text = "\n".join(problems)
    for secret in ("SECRET", "PRIVATE"):
        assert secret not in text


def test_a_workbook_with_unreadable_cells_reports_them_with_the_other_problems(
    tmp_path: Path,
) -> None:
    path = tmp_path / "plan.xlsx"
    write_plan(path, _plan([_item(2), _item(3), _item(4)], usd_to_pen=None))
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    items = workbook["Gastos fijos"]
    items["I3"] = "unos cien"  # expected_amount, row 3
    items["F4"] = "x"  # typical_amount, row 4
    workbook.save(path)
    problems: list[str] = []

    plan = read_plan(path, problems)
    _, more = _check(plan)

    assert problems == [
        "Gastos fijos: expected_amount in row 3 is not a number",
        "Gastos fijos: typical_amount in row 4 is not a number",
    ]
    assert more == ["Meta: usd_to_pen is required"]
    assert [item.row for item in plan.items] == [2]


def test_read_plan_without_a_collector_still_raises_on_the_first_bad_cell(
    tmp_path: Path,
) -> None:
    path = tmp_path / "plan.xlsx"
    write_plan(path, _plan([_item(2)]))
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    workbook["Gastos fijos"]["I2"] = "mucho"
    workbook.save(path)

    with pytest.raises(ValueError, match="expected_amount in row 2"):
        read_plan(path)


def test_the_row_numbers_come_from_the_sheet(tmp_path: Path) -> None:
    path = tmp_path / "plan.xlsx"
    write_plan(path, _plan([_item(2), _item(3)]))

    plan = read_plan(path)

    assert [item.row for item in plan.items] == [2, 3]
