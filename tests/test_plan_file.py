"""`forecasting.plan_file`: the plan workbook's layout and what a re-export keeps. All
descriptions are synthetic (T56, ADR 0048)."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from forecasting.fixed_expenses import Candidate
from forecasting.plan_file import (
    ITEM_COLUMNS,
    META_DEFAULTS,
    PlanFile,
    PlanItem,
    merge,
    read_plan,
    write_plan,
)


def _candidate(
    description: str, amount: float = 100.0, kind: str = "fixed"
) -> Candidate:
    return Candidate("BCP", description, "PEN", "Servicios", 6, amount, kind)


def _cells(path: Path) -> dict[str, list[list[object]]]:
    workbook = load_workbook(path)
    return {
        sheet.title: [[c.value for c in row] for row in sheet.iter_rows()]
        for sheet in workbook
    }


def test_first_export_prefills_the_choice_with_the_proposal() -> None:
    plan = merge(
        [
            _candidate("PLANTED RENT", 1200.0),
            _candidate("PLANTED SHOP", 80.0, "variable"),
        ],
        None,
    )

    assert [(i.description, i.kind, i.expected_amount) for i in plan.items] == [
        ("PLANTED RENT", "fixed", 1200.0),
        ("PLANTED SHOP", "variable", 80.0),
    ]


def test_first_export_meta_has_the_defaults_and_leaves_the_goal_empty() -> None:
    plan = merge([], None)

    assert plan.meta["emergency_months"] == 6
    assert plan.meta["emergency_basis"] == "all"
    assert plan.meta["emergency_account"] == "Ripley"
    assert plan.meta["goal_amount"] is None
    assert plan.meta["usd_to_pen"] is None
    assert set(plan.meta) == set(META_DEFAULTS)


def test_the_workbook_has_the_three_sheets_in_order_and_english_headers(
    tmp_path: Path,
) -> None:
    path = tmp_path / "plan.xlsx"

    write_plan(path, merge([_candidate("PLANTED RENT")], None))

    cells = _cells(path)
    assert list(cells) == ["Instrucciones", "Gastos fijos", "Meta"]
    assert tuple(cells["Gastos fijos"][0]) == ITEM_COLUMNS
    assert [row[0] for row in cells["Meta"][1:]] == list(META_DEFAULTS)
    assert cells["Meta"][0][:2] == ["field", "value"]


def test_the_instructions_are_in_spanish_and_name_the_three_kinds(
    tmp_path: Path,
) -> None:
    path = tmp_path / "plan.xlsx"

    write_plan(path, merge([], None))

    text = " ".join(str(r[0]) for r in _cells(path)["Instrucciones"] if r[0])
    for word in ("fixed", "variable", "ignore", "expected_amount", "make export-plan"):
        assert word in text
    assert "gasto" in text.lower()


def test_write_then_read_gives_the_same_plan(tmp_path: Path) -> None:
    path = tmp_path / "plan.xlsx"
    plan = merge([_candidate("PLANTED RENT", 1200.0)], None)

    write_plan(path, plan)

    assert read_plan(path) == plan


def test_a_re_export_keeps_the_owners_choice_and_refreshes_untouched_rows(
    tmp_path: Path,
) -> None:
    path = tmp_path / "plan.xlsx"
    first = merge(
        [_candidate("PLANTED RENT", 1200.0), _candidate("PLANTED GYM", 90.0)], None
    )
    edited = PlanFile(
        items=tuple(
            PlanItem(**{**i.__dict__, "kind": "ignore", "expected_amount": 0.0})
            if i.description == "PLANTED RENT"
            else i
            for i in first.items
        ),
        meta={**first.meta, "goal_amount": 15000, "usd_to_pen": 3.75},
    )
    write_plan(path, edited)

    again = merge(
        [_candidate("PLANTED RENT", 1250.0), _candidate("PLANTED GYM", 95.0)],
        read_plan(path),
    )

    by_name = {i.description: i for i in again.items}
    assert (by_name["PLANTED RENT"].kind, by_name["PLANTED RENT"].expected_amount) == (
        "ignore",
        0.0,
    )
    assert by_name["PLANTED RENT"].typical_amount == 1250.0
    assert (by_name["PLANTED GYM"].kind, by_name["PLANTED GYM"].expected_amount) == (
        "fixed",
        95.0,
    )
    assert again.meta["goal_amount"] == 15000
    assert again.meta["usd_to_pen"] == 3.75


def test_a_changed_expected_amount_is_an_owner_choice_too() -> None:
    first = merge([_candidate("PLANTED GYM", 90.0)], None)
    edited = PlanFile(
        items=(PlanItem(**{**first.items[0].__dict__, "expected_amount": 99.0}),),
        meta=first.meta,
    )

    again = merge([_candidate("PLANTED GYM", 95.0)], edited)

    assert again.items[0].expected_amount == 99.0
    assert again.items[0].kind == "fixed"


def test_a_row_no_longer_detected_stays_and_says_so() -> None:
    first = merge([_candidate("PLANTED GONE"), _candidate("PLANTED KEPT")], None)

    again = merge([_candidate("PLANTED KEPT")], first)

    by_name = {i.description: i for i in again.items}
    assert "no longer" in by_name["PLANTED GONE"].note
    assert by_name["PLANTED KEPT"].note == ""


def test_a_new_candidate_is_appended_after_the_existing_rows() -> None:
    first = merge([_candidate("PLANTED OLD")], None)

    again = merge([_candidate("PLANTED NEW", 500.0), _candidate("PLANTED OLD")], first)

    assert [i.description for i in again.items] == ["PLANTED OLD", "PLANTED NEW"]


def test_exporting_twice_leaves_the_workbook_identical(tmp_path: Path) -> None:
    path = tmp_path / "plan.xlsx"
    candidates = [
        _candidate("PLANTED RENT", 1200.0),
        _candidate("PLANTED SHOP", 80.0, "variable"),
    ]
    write_plan(path, merge(candidates, None))
    first = _cells(path)

    write_plan(path, merge(candidates, read_plan(path)))

    assert _cells(path) == first


def test_a_file_without_the_expected_sheets_is_rejected(tmp_path: Path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "other.xlsx"
    Workbook().save(path)

    with pytest.raises(ValueError, match="Gastos fijos"):
        read_plan(path)


def test_read_plan_names_the_row_when_expected_amount_is_not_a_number(
    tmp_path: Path,
) -> None:
    path = tmp_path / "plan.xlsx"
    write_plan(path, merge([_candidate("NETFLIX.COM")], None))
    workbook = load_workbook(path)
    workbook["Gastos fijos"]["I2"] = "3,75"
    workbook.save(path)

    with pytest.raises(ValueError, match=r"expected_amount.*row 2"):
        read_plan(path)


def test_a_row_the_owner_added_by_hand_is_kept_with_their_own_note() -> None:
    plan = merge([_candidate("NETFLIX.COM")], None)
    mine = PlanItem("BCP", "GYM", "PEN", "Deporte", 0, 0.0, "", "fixed", 80.0, "mío")

    again = merge([_candidate("NETFLIX.COM")], PlanFile((*plan.items, mine), plan.meta))

    assert again.items[-1] == mine


def test_a_description_that_looks_like_a_formula_survives_the_round_trip(
    tmp_path: Path,
) -> None:
    path = tmp_path / "plan.xlsx"
    plan = merge([_candidate("=SUM(A1)")], None)
    write_plan(path, plan)

    assert [i.description for i in read_plan(path).items] == ["=SUM(A1)"]


def test_a_bad_system_cell_is_reported_by_row_without_echoing_its_value(
    tmp_path: Path,
) -> None:
    path = tmp_path / "plan.xlsx"
    write_plan(path, merge([_candidate("NETFLIX.COM")], None))
    workbook = load_workbook(path)
    workbook["Gastos fijos"]["F2"] = "secret-12,5"
    workbook.save(path)

    with pytest.raises(ValueError, match=r"row 2") as raised:
        read_plan(path)
    assert "secret" not in str(raised.value)
