"""Tests for categorization.labels: the labeling file's shape (T51, ADR 0043)."""

from pathlib import Path

from openpyxl import Workbook, load_workbook

from categorization.labels import (
    CATEGORIES,
    COLUMNS,
    SHEET,
    read_completed,
    write_template,
)


def test_the_template_has_one_row_per_group_with_the_suggestion_prefilled(
    tmp_path: Path,
) -> None:
    path = tmp_path / "f.xlsx"
    write_template(
        path,
        [
            ("BCP", "NETFLIX.COM", 3, "Servicios"),
            ("BCP", "XYZ", 1, "Sin categorizar"),
        ],
    )

    sheet = load_workbook(path)[SHEET]
    header = tuple(c.value for c in sheet[1])
    rows = list(sheet.iter_rows(min_row=2, values_only=True))

    assert header == COLUMNS
    assert rows[0] == ("BCP", "NETFLIX.COM", 3, "Servicios", "Servicios")
    assert rows[1] == ("BCP", "XYZ", 1, "Sin categorizar", "Gastos varios")


def test_a_filled_in_file_reads_back_bank_description_and_the_chosen_category(
    tmp_path: Path,
) -> None:
    path = tmp_path / "f.xlsx"
    write_template(path, [("BCP", "NETFLIX.COM", 3, "Sin categorizar")])
    workbook = load_workbook(path)
    workbook[SHEET].cell(row=2, column=5, value="Servicios")  # the owner's correction
    workbook.save(path)

    rows, problems = read_completed(path)

    assert problems == []
    assert rows == [("BCP", "NETFLIX.COM", "Servicios")]


def test_a_category_not_on_the_list_is_a_problem_naming_the_row_not_the_value(
    tmp_path: Path,
) -> None:
    path = tmp_path / "f.xlsx"
    write_template(path, [("BCP", "NETFLIX.COM", 1, "Servicios")])
    workbook = load_workbook(path)
    workbook[SHEET].cell(row=2, column=5, value="Comida")  # not one of CATEGORIES
    workbook.save(path)

    rows, problems = read_completed(path)

    assert rows == []
    assert problems == ["row 2: category is not one of the list"]
    assert "Comida" not in problems[0]


def test_a_row_with_an_empty_bank_or_description_is_a_problem(tmp_path: Path) -> None:
    path = tmp_path / "f.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = SHEET
    sheet.append(list(COLUMNS))
    sheet.append(["", "NETFLIX.COM", 1, "Servicios", "Servicios"])
    workbook.save(path)

    rows, problems = read_completed(path)

    assert rows == []
    assert problems == ["row 2: bank or description is empty"]


def test_empty_rows_are_ignored(tmp_path: Path) -> None:
    path = tmp_path / "f.xlsx"
    write_template(path, [("BCP", "NETFLIX.COM", 1, "Servicios")])
    workbook = load_workbook(path)
    workbook[SHEET].append([None, None, None, None, None])
    workbook.save(path)

    rows, problems = read_completed(path)

    assert problems == [] and len(rows) == 1


def test_the_wrong_columns_are_refused(tmp_path: Path) -> None:
    path = tmp_path / "f.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = SHEET
    sheet.append(["bank", "description"])
    workbook.save(path)

    rows, problems = read_completed(path)

    assert rows == [] and "the columns are not the template's" in problems[0]


def test_a_missing_sheet_is_refused(tmp_path: Path) -> None:
    path = tmp_path / "f.xlsx"
    Workbook().save(path)

    rows, problems = read_completed(path)

    assert rows == [] and f"no '{SHEET}' sheet" in problems[0]


def test_not_a_workbook_is_refused(tmp_path: Path) -> None:
    path = tmp_path / "f.txt"
    path.write_bytes(b"not a workbook")

    rows, problems = read_completed(path)

    assert rows == [] and "not a readable .xlsx workbook" in problems[0]


def test_categories_are_a_short_fixed_list_including_a_catch_all() -> None:
    assert 5 <= len(CATEGORIES) <= 20
    assert "Gastos varios" in CATEGORIES
    assert len(CATEGORIES) == len(set(CATEGORIES))


def test_an_unmatched_guess_defaults_the_editable_column_to_gastos_varios_not_itself(
    tmp_path: Path,
) -> None:
    """'Sin categorizar' is not on CATEGORIES: left untouched, the file must still
    import cleanly."""
    path = tmp_path / "f.xlsx"
    write_template(path, [("BCP", "UNKNOWN MERCHANT", 1, "Sin categorizar")])

    rows, problems = read_completed(path)

    assert problems == []
    assert rows == [("BCP", "UNKNOWN MERCHANT", "Gastos varios")]
