"""`scripts/export_plan.py`: the query and the CLI wiring, with a fake psycopg
connection (the real one is exercised on the owner's machine). No description here is
ever real (T56, ADR 0048)."""

import stat
from dataclasses import replace
from datetime import date
from pathlib import Path
from typing import Any

import psycopg
import pytest
from openpyxl import load_workbook

from forecasting.fixed_expenses import MonthlySpend
from forecasting.plan_file import PlanFile, read_plan, write_plan
from scripts import export_plan as ep

MONTHS = [date(2026, m, 1) for m in range(1, 7)]
NETFLIX = [("BCP", "NETFLIX.COM", "PEN", "Servicios", m, 44.9) for m in MONTHS]
GROCERIES = [
    ("BCP", "SUPERMERCADO", "PEN", "Alimentación", m, a)
    for m, a in zip(MONTHS, [210.0, 480.0, 95.0, 330.0, 150.0, 400.0], strict=True)
]


class _FakeCursor:
    def __init__(self, rows: list[tuple[Any, ...]]) -> None:
        self._rows = rows
        self.query = ""
        self.params: dict[str, Any] = {}

    def __enter__(self) -> "_FakeCursor":
        return self

    def __exit__(self, *exc: object) -> None:
        return None

    def execute(self, query: str, params: dict[str, Any]) -> None:
        self.query, self.params = query, params

    def fetchall(self) -> list[tuple[Any, ...]]:
        return self._rows


class _FakeConnection:
    def __init__(self, rows: list[tuple[Any, ...]]) -> None:
        self.cursor_ = _FakeCursor(rows)

    def __enter__(self) -> "_FakeConnection":
        return self

    def __exit__(self, *exc: object) -> None:
        return None

    def cursor(self) -> _FakeCursor:
        return self.cursor_


@pytest.fixture
def environment(monkeypatch: pytest.MonkeyPatch) -> None:
    for name, value in {
        "PFP_PG_DATABASE": "pfp",
        "PFP_PG_USER": "pfp",
        "PFP_PG_PASSWORD": "secret",
    }.items():
        monkeypatch.setenv(name, value)


def test_fetch_monthly_spend_is_scoped_to_the_user_and_to_real_spending(
    environment: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake = _FakeConnection(NETFLIX[:1])
    seen: list[str] = []

    def connect(conninfo: str) -> _FakeConnection:
        seen.append(conninfo)
        return fake

    monkeypatch.setattr(psycopg, "connect", connect)

    rows = ep.fetch_monthly_spend("piero")

    assert [(r.bank, r.description, r.currency, r.amount) for r in rows] == [
        ("BCP", "NETFLIX.COM", "PEN", 44.9)
    ]
    assert fake.cursor_.params == {"user": "piero"}
    query = fake.cursor_.query
    assert "gold.rpt_movements" in query
    assert "not is_internal_transfer" in query
    assert "flow_type = 'egreso'" in query
    assert "read_only" in seen[0] or "read-only" in seen[0].replace("_", "-")


def test_main_writes_a_private_workbook_and_prints_only_counts(
    environment: None,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    rows = [MonthlySpend(*r) for r in NETFLIX + GROCERIES]
    monkeypatch.setattr(ep, "fetch_monthly_spend", lambda user_id: rows)
    out = tmp_path / "plan" / "plan-de-ahorro.xlsx"

    assert ep.main(["--user", "piero", "--out", str(out)]) == 0

    assert load_workbook(out).sheetnames == ["Instrucciones", "Gastos fijos", "Meta"]
    assert stat.S_IMODE(out.stat().st_mode) == 0o600
    assert stat.S_IMODE(out.parent.stat().st_mode) == 0o700
    printed = capsys.readouterr().out
    assert "2 candidate(s)" in printed
    assert "1 proposed as fixed" in printed
    assert "NETFLIX" not in printed
    assert "44.9" not in printed


def test_main_keeps_the_owner_choices_on_a_second_export(
    environment: None, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    rows = [MonthlySpend(*r) for r in NETFLIX + GROCERIES]
    monkeypatch.setattr(ep, "fetch_monthly_spend", lambda user_id: rows)
    out = tmp_path / "plan.xlsx"
    ep.main(["--user", "piero", "--out", str(out)])
    first = read_plan(out)
    edited = [
        replace(item, kind="ignore") if item.description == "NETFLIX.COM" else item
        for item in first.items
    ]
    write_plan(out, PlanFile(tuple(edited), {**first.meta, "usd_to_pen": 3.7}))

    assert ep.main(["--user", "piero", "--out", str(out)]) == 0

    again = read_plan(out)
    kinds = {item.description: item.kind for item in again.items}
    assert kinds["NETFLIX.COM"] == "ignore"
    assert again.meta["usd_to_pen"] == 3.7


def test_main_refuses_without_a_user(
    environment: None, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    monkeypatch.delenv("PFP_USER", raising=False)
    assert ep.main(["--out", str(tmp_path / "p.xlsx")]) == 2


def test_main_says_so_when_gold_has_no_spending(
    environment: None, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    monkeypatch.setattr(ep, "fetch_monthly_spend", lambda user_id: [])
    out = tmp_path / "p.xlsx"
    assert ep.main(["--user", "piero", "--out", str(out)]) == 1
    assert not out.exists()


def test_the_staging_file_is_private_and_removed_when_writing_fails(
    environment: None, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    rows = [MonthlySpend(*r) for r in NETFLIX]
    monkeypatch.setattr(ep, "fetch_monthly_spend", lambda user_id: rows)
    seen: dict[str, int] = {}
    real_write = ep.write_plan

    def failing_write(path: Path, plan: Any) -> None:
        real_write(path, plan)
        seen["mode"] = stat.S_IMODE(path.stat().st_mode)
        raise OSError("disk full")

    monkeypatch.setattr(ep, "write_plan", failing_write)
    out = tmp_path / "plan.xlsx"

    with pytest.raises(OSError):
        ep.main(["--user", "piero", "--out", str(out)])

    assert seen["mode"] == 0o600
    assert list(tmp_path.iterdir()) == []
