"""The manual Excel's layout, the generic template and the structure check the upload
portal runs without any secret (T49, ADR 0041)."""

import io
import zipfile
from datetime import date

from openpyxl import Workbook, load_workbook

from ingestion import manual_layout as layout


def _workbook(
    savings: list[tuple[object, ...]] | None = None,
    funds: list[tuple[object, ...]] | None = None,
    *,
    savings_columns: tuple[str, ...] = layout.SAVINGS_COLUMNS,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = layout.SHEET
    sheet.append(list(savings_columns))
    for row in savings or []:
        sheet.append(list(row))
    second = workbook.create_sheet(layout.INVESTMENT_SHEET)
    second.append(list(layout.INVESTMENT_COLUMNS))
    for row in funds or []:
        second.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


_FUND = ("Mi fondo", date(2026, 7, 5), "aporte", 100, "PEN", 100, None)


def test_the_generic_template_has_both_sheets_free_funds_and_only_example_rows() -> (
    None
):
    workbook = load_workbook(io.BytesIO(layout.generic_template()))

    assert workbook.sheetnames == ["Ahorros", "Inversiones", "Instrucciones"]
    funds = workbook["Inversiones"]
    assert [c.value for c in funds[1]][: len(layout.INVESTMENT_COLUMNS)] == list(
        layout.INVESTMENT_COLUMNS
    )
    assert all(
        str(r[0]).startswith("EJEMPLO")
        for r in funds.iter_rows(min_row=2, values_only=True)
    )
    # No drop-down on the fund's name, unlike the owner's own template.
    assert not any("A" in str(v.sqref) for v in funds.data_validations.dataValidation)


def test_the_untouched_template_is_refused_because_of_its_example_rows() -> None:
    problems = layout.check_workbook(layout.generic_template())

    assert any("example row(s) are still in the sheet" in p for p in problems)


def test_a_workbook_with_one_sheet_filled_and_the_other_empty_passes() -> None:
    assert layout.check_workbook(_workbook(funds=[_FUND])) == []


def test_a_workbook_with_both_sheets_empty_is_refused() -> None:
    assert layout.check_workbook(_workbook()) == [
        "both sheets are empty: fill in at least one"
    ]


def test_changed_columns_or_a_missing_sheet_are_named() -> None:
    wrong = _workbook(funds=[_FUND], savings_columns=("cuenta", "otra"))

    assert layout.check_workbook(wrong) == [
        "Ahorros: the columns are not the template's"
    ]
    lone = Workbook()
    lone.save(buffer := io.BytesIO())
    assert "the workbook has no 'Ahorros' sheet" in layout.check_workbook(
        buffer.getvalue()
    )


def test_something_that_is_not_a_workbook_is_refused() -> None:
    assert layout.check_workbook(b"not a workbook") == [
        "the file is not a readable .xlsx workbook"
    ]


def test_a_workbook_with_macros_is_refused() -> None:
    buffer = io.BytesIO(_workbook(funds=[_FUND]))
    with zipfile.ZipFile(buffer, "a") as archive:
        archive.writestr("xl/vbaProject.bin", b"macro")

    assert layout.check_workbook(buffer.getvalue()) == [
        "the workbook has macros, which are not accepted"
    ]


def test_a_zip_that_unpacks_to_far_more_than_a_workbook_is_refused() -> None:
    buffer = io.BytesIO(_workbook(funds=[_FUND]))
    with zipfile.ZipFile(buffer, "a", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/big.bin", b"0" * (layout.MAX_UNCOMPRESSED_BYTES + 1))

    problems = layout.check_workbook(buffer.getvalue())

    assert problems == ["the file is far bigger than a workbook of this kind"]


def test_a_file_over_five_megabytes_is_refused() -> None:
    assert layout.check_workbook(b"0" * (layout.MAX_BYTES + 1)) == [
        "the file is larger than 5 MB"
    ]


def test_the_problems_never_quote_a_cell() -> None:
    secret = "Cuenta Con Numero 191-4827"
    problems = layout.check_workbook(
        _workbook(savings=[("EJEMPLO " + secret, date(2026, 1, 1), "x", 1, "PEN", 1)])
    )

    assert problems and not any(secret in p for p in problems)
