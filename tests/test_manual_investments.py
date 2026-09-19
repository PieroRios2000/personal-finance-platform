"""Reading the manual Excel's `Inversiones` sheet. Everything here is invented;
problems are row numbers and column names, never values."""

from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

from ingestion.manual_excel import INVESTMENT_COLUMNS, read_investments
from scripts.make_manual_templates import write_template


def _workbook(path: Path, rows: list[tuple[Any, ...]] | None) -> Path:
    workbook = Workbook()
    savings = workbook.active
    assert savings is not None
    savings.title = "Ahorros"
    if rows is not None:
        sheet = workbook.create_sheet("Inversiones")
        sheet.append(list(INVESTMENT_COLUMNS))
        for row in rows:
            sheet.append(list(row))
    workbook.save(path)
    return path


def _row(
    day: date,
    kind: str,
    amount: float,
    balance: float,
    place: str = "Fondo A",
    currency: str = "PEN",
    note: str | None = None,
) -> tuple[Any, ...]:
    return (place, day, kind, amount, currency, balance, note)


_TWO_MONTHS = [
    _row(date(2026, 7, 5), "aporte", 100, 100, note="Nombre Real Del Fondo"),
    _row(date(2026, 7, 31), "valorizacion", 0, 102),
    _row(date(2026, 8, 10), "retiro", 50, 53),
    _row(date(2026, 8, 31), "valorizacion", 0, 54),
]


def test_each_fund_month_becomes_one_bronze_month_with_its_rows(tmp_path: Path) -> None:
    result = read_investments(
        _workbook(tmp_path / "f.xlsx", _TWO_MONTHS), user_id="piero"
    )

    assert result.problems == []
    july, august = result.months
    assert (july.place, july.currency, july.year, july.month) == (
        "Fondo A",
        "PEN",
        2026,
        7,
    )
    assert [(e.kind, e.amount, e.balance) for e in july.entries] == [
        ("aporte", Decimal("100.00"), Decimal("100.00")),
        ("valorizacion", Decimal("0.00"), Decimal("102.00")),
    ]
    assert [e.position for e in august.entries] == sorted(
        e.position for e in august.entries
    )


def test_the_note_is_kept_as_the_funds_detail(tmp_path: Path) -> None:
    result = read_investments(
        _workbook(tmp_path / "f.xlsx", _TWO_MONTHS), user_id="piero"
    )

    assert result.months[0].entries[0].detail == "Nombre Real Del Fondo"
    assert result.months[0].entries[1].detail is None


def test_a_months_identity_does_not_depend_on_its_content(tmp_path: Path) -> None:
    first = read_investments(
        _workbook(tmp_path / "a.xlsx", _TWO_MONTHS), user_id="piero"
    )
    changed = [_row(date(2026, 7, 6), "aporte", 200, 200), *_TWO_MONTHS[1:]]
    second = read_investments(_workbook(tmp_path / "b.xlsx", changed), user_id="piero")

    assert first.months[0].month_key == second.months[0].month_key
    assert first.months[0].month_key != first.months[1].month_key


def test_funds_and_currencies_are_separate_months(tmp_path: Path) -> None:
    rows = [
        _row(date(2026, 7, 5), "aporte", 100, 100, place="Fondo A"),
        _row(date(2026, 7, 5), "aporte", 100, 100, place="Fondo B"),
        _row(date(2026, 7, 5), "aporte", 10, 10, place="Fondo A", currency="USD"),
    ]

    result = read_investments(_workbook(tmp_path / "f.xlsx", rows), user_id="piero")

    assert len({m.month_key for m in result.months}) == 3


def test_months_without_a_valuation_and_missing_months_are_counted(
    tmp_path: Path,
) -> None:
    rows = [
        _row(date(2026, 7, 5), "aporte", 100, 100),
        _row(date(2026, 9, 5), "aporte", 100, 205),
        _row(date(2026, 9, 30), "valorizacion", 0, 206),
    ]

    result = read_investments(_workbook(tmp_path / "f.xlsx", rows), user_id="piero")

    assert result.months_without_valuation == 1
    assert result.missing_months == 1


@pytest.mark.parametrize(
    ("row", "expected"),
    [
        (("", date(2026, 7, 5), "aporte", 1, "PEN", 1, None), "lugar is empty"),
        (("F", "no", "aporte", 1, "PEN", 1, None), "fecha is not a date"),
        (
            ("F", date(2026, 7, 5), "compra", 1, "PEN", 1, None),
            "tipo must be aporte, retiro or valorizacion",
        ),
        (
            ("F", date(2026, 7, 5), "aporte", "x", "PEN", 1, None),
            "monto is not a number",
        ),
        (
            ("F", date(2026, 7, 5), "aporte", 0, "PEN", 1, None),
            "monto must be greater than 0 in an aporte or retiro",
        ),
        (
            ("F", date(2026, 7, 5), "valorizacion", 5, "PEN", 1, None),
            "monto must be 0 in a valorizacion",
        ),
        (
            ("F", date(2026, 7, 5), "aporte", 1, "EUR", 1, None),
            "moneda must be PEN or USD",
        ),
        (
            ("F", date(2026, 7, 5), "aporte", 1, "PEN", None, None),
            "saldo_final is empty (a formula without a saved value? "
            "open and save the file in Excel)",
        ),
        (
            ("F", date(2026, 7, 5), "aporte", 1, "PEN", -5, None),
            "saldo_final cannot be negative",
        ),
    ],
)
def test_invalid_cells_are_reported_by_row_and_column_never_by_value(
    tmp_path: Path, row: tuple[Any, ...], expected: str
) -> None:
    result = read_investments(_workbook(tmp_path / "f.xlsx", [row]), user_id="piero")

    assert result.months == []
    assert f"Inversiones row 2: {expected}" in result.problems


def test_a_workbook_without_the_sheet_or_with_no_rows_imports_nothing_quietly(
    tmp_path: Path,
) -> None:
    for name, rows in (("a", None), ("b", [])):
        result = read_investments(
            _workbook(tmp_path / f"{name}.xlsx", rows), user_id="piero"
        )

        assert result.months == [] and result.problems == []


def test_example_rows_left_in_the_template_are_a_problem(tmp_path: Path) -> None:
    template = tmp_path / "t.xlsx"
    write_template(template)

    result = read_investments(template, user_id="piero")

    assert result.months == []
    assert result.problems == [
        "Inversiones: 4 example row(s) are still in the sheet (EJEMPLO)"
    ]


def test_different_columns_are_a_clear_problem(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Inversiones"
    sheet.append(["lugar", "otra"])
    workbook.save(tmp_path / "f.xlsx")

    assert read_investments(tmp_path / "f.xlsx", user_id="piero").problems == [
        "Inversiones: the columns are not the template's"
    ]
