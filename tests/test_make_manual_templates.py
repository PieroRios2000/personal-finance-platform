"""The Excel template for data the banks won't export as statements: savings
account movements (Ripley) and investments, as two sheets of one workbook. Every
row in it is invented."""

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ingestion.manual_excel import SAVINGS_COLUMNS
from scripts.make_manual_templates import INVESTMENT_COLUMNS, main, write_template


def _sheet(path: Path, name: str) -> Worksheet:
    return load_workbook(path)[name]


def _text(sheet: Worksheet) -> str:
    return " ".join(str(c.value) for row in sheet.iter_rows() for c in row if c.value)


def _rows(sheet: Worksheet) -> list[tuple[Any, ...]]:
    return list(sheet.iter_rows(min_row=2, values_only=True))


def test_one_workbook_with_a_savings_sheet_an_investments_sheet_and_instructions(
    tmp_path: Path,
) -> None:
    path = tmp_path / "finanzas-manual.xlsx"

    write_template(path)

    assert load_workbook(path).sheetnames == ["Ahorros", "Inversiones", "Instrucciones"]


def test_the_savings_sheet_has_exactly_the_agreed_columns(tmp_path: Path) -> None:
    path = tmp_path / "f.xlsx"
    write_template(path)

    header = [c.value for c in _sheet(path, "Ahorros")[1]]

    assert header == [
        "cuenta",
        "fecha",
        "descripcion",
        "monto",
        "moneda",
        "saldo_final",
    ]
    assert SAVINGS_COLUMNS == tuple(header)


def test_the_investments_sheet_has_exactly_the_agreed_columns(tmp_path: Path) -> None:
    path = tmp_path / "f.xlsx"
    write_template(path)

    header = [c.value for c in _sheet(path, "Inversiones")[1]]

    assert header == [
        "lugar",
        "fecha",
        "tipo",
        "monto",
        "moneda",
        "saldo_final",
        "nota",
    ]
    assert INVESTMENT_COLUMNS == tuple(header)


def test_savings_examples_are_typed_cells_and_obviously_fictional(
    tmp_path: Path,
) -> None:
    path = tmp_path / "f.xlsx"
    write_template(path)

    rows = _rows(_sheet(path, "Ahorros"))

    assert rows, "the template shows how to fill it in"
    for cuenta, fecha, descripcion, monto, moneda, saldo in rows:
        assert str(cuenta).startswith("EJEMPLO")
        assert isinstance(fecha, (date, datetime))
        assert isinstance(monto, (int, float))
        assert isinstance(saldo, (int, float))
        assert moneda in {"PEN", "USD"}
        assert descripcion


def test_savings_examples_show_the_month_end_row_with_a_zero_amount(
    tmp_path: Path,
) -> None:
    path = tmp_path / "f.xlsx"
    write_template(path)

    rows = _rows(_sheet(path, "Ahorros"))

    assert any(r[2] == "cierre de mes" and r[3] == 0 for r in rows)


def test_investment_examples_are_typed_cells_and_obviously_fictional(
    tmp_path: Path,
) -> None:
    path = tmp_path / "f.xlsx"
    write_template(path)

    rows = _rows(_sheet(path, "Inversiones"))

    assert rows
    for _lugar, fecha, tipo, monto, moneda, saldo, nota in rows:
        assert isinstance(fecha, (date, datetime))
        assert isinstance(monto, (int, float))
        assert isinstance(saldo, (int, float))
        assert tipo in {"aporte", "retiro", "valorizacion"}
        assert moneda in {"PEN", "USD"}
        assert str(nota).startswith("EJEMPLO")
    assert any(r[2] == "valorizacion" and r[3] == 0 for r in rows)


def test_the_dropdowns_only_allow_the_agreed_values(tmp_path: Path) -> None:
    path = tmp_path / "f.xlsx"
    write_template(path)

    savings = "".join(
        v.formula1 for v in _sheet(path, "Ahorros").data_validations.dataValidation
    )
    investments = "".join(
        v.formula1 for v in _sheet(path, "Inversiones").data_validations.dataValidation
    )

    assert "PEN,USD" in savings
    assert "aporte,retiro,valorizacion" in investments
    assert "PEN,USD" in investments
    assert "Tyba fondo 1" in investments


def test_the_instructions_explain_every_column_and_the_two_special_rows(
    tmp_path: Path,
) -> None:
    path = tmp_path / "f.xlsx"
    write_template(path)

    text = _text(_sheet(path, "Instrucciones"))

    for column in (*SAVINGS_COLUMNS, *INVESTMENT_COLUMNS):
        assert column in text
    assert "valorizacion" in text
    assert "cierre de mes" in text


def test_main_writes_the_template_privately_and_never_under_the_real_name(
    tmp_path: Path,
) -> None:
    out = tmp_path / "manual"

    assert main(["--out-dir", str(out)]) == 0

    assert [p.name for p in out.iterdir()] == ["plantilla-finanzas-manual.xlsx"]
    assert out.stat().st_mode & 0o777 == 0o700
    assert (out / "plantilla-finanzas-manual.xlsx").stat().st_mode & 0o777 == 0o600
