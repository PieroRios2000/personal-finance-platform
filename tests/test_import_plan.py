"""`scripts/import_plan.py`: validates the plan workbook and calls
`lakehouse.bronze.replace_plan`; the Delta write itself is `tests/test_bronze.py`'s job
(T57, ADR 0048). Every description here is synthetic."""

import stat
from datetime import date
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from forecasting.plan_file import PlanFile, PlanItem, write_plan
from lakehouse import bronze
from scripts import export_plan
from scripts import import_plan as ip


def _item(row: int, **overrides: object) -> PlanItem:
    fields: dict[str, Any] = {
        "bank": "BCP",
        "description": f"PLANTED ITEM {row}",
        "currency": "PEN",
        "category": "Servicios",
        "months_seen": 6,
        "typical_amount": 50.0,
        "proposed_kind": "fixed",
        "kind": "fixed",
        "expected_amount": 50.0,
        "row": row,
    }
    fields.update(overrides)
    return PlanItem(**fields)


GOOD_META: dict[str, object] = {
    "goal_amount": 10000,
    "usd_to_pen": 3.75,
    "emergency_months": 6,
    "emergency_basis": "all",
    "emergency_account": "Ripley",
    "target_date": date(2099, 12, 31),
    "income_pen_override": None,
    "income_usd_override": None,
}


def _workbook(tmp_path: Path, items: list[PlanItem], **meta: object) -> Path:
    path = tmp_path / "plan.xlsx"
    write_plan(path, PlanFile(tuple(items), {**GOOD_META, **meta}))
    return path


@pytest.fixture(autouse=True)
def environment(monkeypatch: pytest.MonkeyPatch) -> list[Any]:
    monkeypatch.setenv("LAKEHOUSE_URI", "s3://lakehouse")
    monkeypatch.setenv("AWS_ENDPOINT_URL", "http://localhost:8333")
    monkeypatch.setattr(
        bronze, "asset_account_names", lambda user_id: {"BCP", "Ripley"}
    )
    calls: list[Any] = []
    monkeypatch.setattr(
        bronze,
        "replace_plan",
        lambda user_id, items, goal: calls.append((user_id, items, goal)),
    )
    return calls


def test_a_good_plan_replaces_the_users_plan(
    environment: list[Any], tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    path = _workbook(
        tmp_path,
        [_item(2), _item(3, kind="variable"), _item(4, kind="ignore")],
    )

    assert ip.main(["--user", "piero", "--workbook", str(path)]) == 0

    [(user_id, items, goal)] = environment
    assert user_id == "piero"
    assert [(i.description, i.kind, i.expected_amount) for i in items] == [
        ("PLANTED ITEM 2", "fixed", 50.0),
        ("PLANTED ITEM 3", "variable", 50.0),
        ("PLANTED ITEM 4", "ignore", 50.0),
    ]
    assert goal.goal_amount == 10000.0
    assert goal.usd_to_pen == 3.75
    assert goal.target_date == date(2099, 12, 31)
    out = capsys.readouterr().out
    assert "3 item(s)" in out
    assert "1 fixed, 1 variable, 1 ignore" in out
    assert "PLANTED" not in out


def test_a_bad_plan_writes_nothing_and_lists_every_problem(
    environment: list[Any], tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    path = _workbook(
        tmp_path,
        [_item(2), _item(3, kind="fijo"), _item(4, expected_amount=None)],
        usd_to_pen=None,
    )
    workbook = load_workbook(path)
    workbook["Gastos fijos"]["I2"] = "unos cien"
    workbook.save(path)

    assert ip.main(["--user", "piero", "--workbook", str(path)]) == 1

    assert environment == []
    err = capsys.readouterr().err
    assert "expected_amount in row 2 is not a number" in err
    assert "kind in row 3 must be fixed, variable or ignore" in err
    assert "expected_amount in row 4 must be a positive number" in err
    assert "Meta: usd_to_pen is required" in err
    assert "nothing was written" in err
    assert "PLANTED" not in err


def test_the_emergency_account_is_checked_against_the_lake(
    environment: list[Any], tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    path = _workbook(tmp_path, [_item(2)], emergency_account="Cajamarca")

    assert ip.main(["--user", "piero", "--workbook", str(path)]) == 1

    assert environment == []
    assert "emergency_account does not match an asset account" in (
        capsys.readouterr().err
    )


def test_a_workbook_missing_a_sheet_is_a_clear_error(
    environment: list[Any], tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    path = _workbook(tmp_path, [_item(2)])
    workbook = load_workbook(path)
    del workbook["Meta"]
    workbook.save(path)

    assert ip.main(["--user", "piero", "--workbook", str(path)]) == 1

    assert environment == []
    assert "has no sheet 'Meta'" in capsys.readouterr().err


def test_a_missing_workbook_is_reported(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    assert ip.main(["--user", "piero", "--workbook", str(tmp_path / "nope.xlsx")]) == 2

    assert "workbook not found" in capsys.readouterr().err


def test_an_unset_lake_is_reported(
    environment: list[Any],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.delenv("LAKEHOUSE_URI")
    path = _workbook(tmp_path, [_item(2)])

    assert ip.main(["--user", "piero", "--workbook", str(path)]) == 1

    assert environment == []
    assert "LAKEHOUSE_URI" in capsys.readouterr().err


def test_a_world_readable_workbook_is_made_private_even_when_rejected(
    environment: list[Any], tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    path = _workbook(tmp_path, [_item(2, kind="x")])
    path.chmod(0o644)

    assert ip.main(["--user", "piero", "--workbook", str(path)]) == 1

    assert stat.S_IMODE(path.stat().st_mode) == 0o600
    assert "made it private" in capsys.readouterr().err


def test_a_private_workbook_prints_no_permission_note(
    environment: list[Any], tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    path = _workbook(tmp_path, [_item(2)])
    path.chmod(0o600)

    assert ip.main(["--user", "piero", "--workbook", str(path)]) == 0

    assert "private" not in capsys.readouterr().err


def test_the_default_workbook_is_the_one_export_plan_writes() -> None:
    assert ip.DEFAULT_WORKBOOK == export_plan.DEFAULT_OUT
