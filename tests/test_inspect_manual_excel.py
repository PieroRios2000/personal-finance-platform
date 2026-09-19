"""The masked description of the owner's manual Excel (ADR 0004): structure and
counts only. Built on a synthetic workbook; the real one is never in the repo."""

from datetime import date
from pathlib import Path

from openpyxl import Workbook

from scripts.inspect_manual_excel import describe, main

_SAVINGS_HEADER = ["cuenta", "fecha", "descripcion", "monto", "moneda", "saldo_final"]
_INVEST_HEADER = ["lugar", "fecha", "tipo", "monto", "moneda", "saldo_final", "nota"]


def _workbook(
    path: Path, savings: list[list[object]], invest: list[list[object]]
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Ahorros"
    sheet.append(_SAVINGS_HEADER)
    for row in savings:
        sheet.append(row)
    second = workbook.create_sheet("Inversiones")
    second.append(_INVEST_HEADER)
    for row in invest:
        second.append(row)
    workbook.save(path)


_SAVINGS = [
    [
        "Mi Cuenta Secreta",
        date(2026, 7, 5),
        "Deposito Juan Perez",
        1234.56,
        "PEN",
        1234.56,
    ],
    ["Mi Cuenta Secreta", date(2026, 7, 20), "Intereses", 7.5, "PEN", 1242.06],
    ["Mi Cuenta Secreta", date(2026, 8, 31), "cierre de mes", 0, "PEN", 1242.06],
    ["Mi Cuenta Secreta", date(2026, 9, 30), "cierre de mes", 0, "PEN", 999.99],
]
_INVEST = [
    [
        "Tyba fondo 1",
        date(2026, 7, 5),
        "aporte",
        500,
        "PEN",
        500,
        "Fondo Renta Ultra Secreto",
    ],
    [
        "Tyba fondo 1",
        date(2026, 7, 31),
        "valorizacion",
        0,
        "PEN",
        505,
        "Fondo Renta Ultra Secreto",
    ],
    ["Tyba fondo 1", date(2026, 8, 31), "valorizacion", 0, "PEN", 510, None],
]


def test_no_real_value_ever_appears(tmp_path: Path) -> None:
    path = tmp_path / "f.xlsx"
    _workbook(path, _SAVINGS, _INVEST)

    text = describe(path)

    for leaked in (
        "Secreta",
        "Juan",
        "Perez",
        "1234",
        "1242",
        "Ultra",
        "Renta",
        "505",
        "999.99",
    ):
        assert leaked not in text
    assert str(path) not in text


def test_the_expected_column_names_are_shown_and_others_are_masked(
    tmp_path: Path,
) -> None:
    path = tmp_path / "f.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Ahorros"
    sheet.append(["cuenta", "fecha", "Numero Secreto 12345"])
    workbook.save(path)

    text = describe(path)

    assert "cuenta" in text
    assert "Numero Secreto" not in text
    assert "12345" not in text


def test_it_reports_sheets_and_row_counts(tmp_path: Path) -> None:
    path = tmp_path / "f.xlsx"
    _workbook(path, _SAVINGS, _INVEST)

    text = describe(path)

    assert "Ahorros" in text and "rows: 4" in text
    assert "Inversiones" in text and "rows: 3" in text


def test_text_values_are_shown_as_shapes_with_counts(tmp_path: Path) -> None:
    path = tmp_path / "f.xlsx"
    _workbook(path, _SAVINGS, _INVEST)

    text = describe(path)

    assert "XX XXXXXX XXXXXXX (4)" in text  # the account name shape


def test_it_counts_how_the_balance_chain_holds(tmp_path: Path) -> None:
    """Each row's balance must be the previous one plus its amount, per account:
    reported as a count, never a value."""
    path = tmp_path / "f.xlsx"
    _workbook(path, _SAVINGS, _INVEST)

    text = describe(path)

    # rows 2 and 3 chain; the last row (999.99) breaks it
    assert "balance chain: 2 of 3 rows follow the previous balance" in text


def test_it_counts_months_and_month_end_markers(tmp_path: Path) -> None:
    path = tmp_path / "f.xlsx"
    _workbook(path, _SAVINGS, _INVEST)

    text = describe(path)

    assert "months covered per group: 3" in text
    assert "cierre de mes rows: 2" in text


def test_investments_report_the_types_and_months_without_a_valuation(
    tmp_path: Path,
) -> None:
    path = tmp_path / "f.xlsx"
    _workbook(path, _SAVINGS, _INVEST)

    text = describe(path)

    assert "tipo: aporte=1, valorizacion=2" in text
    assert "months with a valuation per group: 2 of 2" in text


def test_a_month_without_a_valuation_is_counted(tmp_path: Path) -> None:
    path = tmp_path / "f.xlsx"
    invest = [
        ["Flip", date(2026, 7, 5), "aporte", 100, "PEN", 100, None],
        ["Flip", date(2026, 8, 20), "aporte", 100, "PEN", 200, None],
        ["Flip", date(2026, 8, 31), "valorizacion", 0, "PEN", 201, None],
    ]
    _workbook(path, _SAVINGS, invest)

    assert "months with a valuation per group: 1 of 2" in describe(path)


def test_a_missing_file_is_an_error(tmp_path: Path) -> None:
    assert main([str(tmp_path / "nope.xlsx")]) == 2
