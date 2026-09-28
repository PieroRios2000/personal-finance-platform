"""`scripts/import_category_labels.py`: validates the completed file and calls
`lakehouse.bronze.replace_category_labels` -- that call's own effect (the real Delta
write) is `tests/test_bronze_integration.py`'s job, not this one's (T51, ADR 0043)."""

from pathlib import Path
from typing import Any

import pytest

from categorization.labels import write_template
from lakehouse import bronze
from scripts import import_category_labels as icl


@pytest.fixture(autouse=True)
def environment(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("LAKEHOUSE_URI", "s3://lakehouse")
    monkeypatch.setenv("AWS_ENDPOINT_URL", "http://localhost:8333")


def test_a_completed_file_replaces_the_users_labels(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    path = tmp_path / "labels.xlsx"
    write_template(path, [("BCP", "NETFLIX.COM", 3, "Servicios")])
    calls: list[Any] = []
    monkeypatch.setattr(
        bronze,
        "replace_category_labels",
        lambda user_id, labels: calls.append((user_id, labels)),
    )

    assert icl.main([str(path), "--user", "piero"]) == 0

    # An accepted, non-"Sin categorizar" suggestion left untouched: not trusted.
    assert calls == [("piero", [("BCP", "NETFLIX.COM", "Servicios", False)])]


def test_a_correction_is_trusted_and_reported_separately(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    from openpyxl import load_workbook

    path = tmp_path / "labels.xlsx"
    write_template(
        path,
        [
            ("BCP", "NETFLIX.COM", 3, "Servicios"),  # accepted as is
            ("BCP", "XYZ CORP", 1, "Sin categorizar"),  # no rule matched
        ],
    )
    workbook = load_workbook(path)
    workbook["Categorias"].cell(row=3, column=5, value="Restaurantes")  # corrected
    workbook.save(path)
    monkeypatch.setattr(bronze, "replace_category_labels", lambda user_id, labels: None)

    assert icl.main([str(path), "--user", "piero"]) == 0

    out = capsys.readouterr().out
    assert "1 reviewed" in out
    assert "1 accepted the guesser's suggestion as is" in out


def test_a_problem_writes_nothing(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    path = tmp_path / "labels.xlsx"
    write_template(path, [("BCP", "NETFLIX.COM", 3, "Servicios")])
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    workbook["Categorias"].cell(row=2, column=5, value="Comida")  # not on the list
    workbook.save(path)
    calls: list[Any] = []
    monkeypatch.setattr(
        bronze, "replace_category_labels", lambda user_id, labels: calls.append(labels)
    )

    assert icl.main([str(path), "--user", "piero"]) == 1

    assert calls == []


def test_a_missing_workbook_is_refused() -> None:
    assert icl.main(["/tmp/does-not-exist.xlsx", "--user", "piero"]) == 2


def test_requires_the_lakehouse_uri(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    monkeypatch.delenv("LAKEHOUSE_URI", raising=False)
    path = tmp_path / "labels.xlsx"
    write_template(path, [("BCP", "NETFLIX.COM", 1, "Servicios")])

    assert icl.main([str(path), "--user", "piero"]) == 1
