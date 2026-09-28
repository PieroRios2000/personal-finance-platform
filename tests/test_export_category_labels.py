"""`scripts/export_category_labels.py`: the query and the CLI wiring, with a fake
psycopg connection -- the real connection is what `dbt build`/CI's ephemeral postgres
exercises. No description here is ever real (T51, ADR 0043)."""

from pathlib import Path
from typing import Any

import psycopg
import pytest
from openpyxl import load_workbook

from scripts import export_category_labels as ecl


class _FakeCursor:
    def __init__(self, rows: list[tuple[str, str, int]]) -> None:
        self._rows = rows
        self.seen_params: dict[str, Any] = {}

    def __enter__(self) -> "_FakeCursor":
        return self

    def __exit__(self, *exc: object) -> None:
        return None

    def execute(self, query: str, params: dict[str, Any]) -> None:
        self.seen_params = params

    def fetchall(self) -> list[tuple[str, str, int]]:
        return self._rows


class _FakeConnection:
    def __init__(self, rows: list[tuple[str, str, int]]) -> None:
        self._cursor = _FakeCursor(rows)

    def __enter__(self) -> "_FakeConnection":
        return self

    def __exit__(self, *exc: object) -> None:
        return None

    def cursor(self) -> _FakeCursor:
        return self._cursor


@pytest.fixture
def environment(monkeypatch: pytest.MonkeyPatch) -> None:
    for name, value in {
        "PFP_PG_DATABASE": "pfp",
        "PFP_PG_USER": "pfp",
        "PFP_PG_PASSWORD": "secret",
    }.items():
        monkeypatch.setenv(name, value)


def test_fetch_groups_queries_gold_scoped_to_the_user_excluding_transfers(
    environment: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake = _FakeConnection([("BCP", "NETFLIX.COM", 3)])
    monkeypatch.setattr(psycopg, "connect", lambda conninfo: fake)

    groups = ecl.fetch_groups("piero")

    assert groups == [("BCP", "NETFLIX.COM", 3)]
    assert fake._cursor.seen_params == {"user": "piero"}


def test_main_writes_the_template_with_a_guess_already_filled_in(
    environment: None, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    monkeypatch.setattr(
        ecl,
        "fetch_groups",
        lambda user_id: [("BCP", "NETFLIX.COM", 3), ("BCP", "XYZ CORP", 1)],
    )
    out = tmp_path / "labels.xlsx"

    assert (
        ecl.main(
            [
                "--user",
                "piero",
                "--out",
                str(out),
                "--model-path",
                str(tmp_path / "no-model-here.joblib"),
            ]
        )
        == 0
    )

    rows = list(load_workbook(out)["Categorias"].iter_rows(min_row=2, values_only=True))
    assert rows[0] == ("BCP", "NETFLIX.COM", 3, "Servicios", "Servicios")
    assert rows[1] == ("BCP", "XYZ CORP", 1, "Sin categorizar", "Gastos varios")


def test_main_uses_the_trained_model_when_one_has_been_saved(
    environment: None, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    monkeypatch.setattr(
        ecl, "fetch_groups", lambda user_id: [("BCP", "SOME NEW MERCHANT", 2)]
    )
    model_path = tmp_path / "model.joblib"
    model_path.write_bytes(b"stand-in; ecl.load_model is monkeypatched below")
    monkeypatch.setattr(ecl, "load_model", lambda path: object())
    monkeypatch.setattr(
        ecl, "suggest_from_model", lambda bundle, description, guess: "Deporte"
    )
    out = tmp_path / "labels.xlsx"

    assert (
        ecl.main(
            ["--user", "piero", "--out", str(out), "--model-path", str(model_path)]
        )
        == 0
    )

    rows = list(load_workbook(out)["Categorias"].iter_rows(min_row=2, values_only=True))
    assert rows[0] == ("BCP", "SOME NEW MERCHANT", 2, "Deporte", "Deporte")


def test_main_finds_the_model_even_when_model_path_is_given_without_joblib(
    environment: None, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """train_category_model.py's own --model-path/PFP_CATEGORY_MODEL_PATH convention
    accepts a path with no suffix and appends ".joblib" itself when saving; this
    script has to resolve the exact same way, or pointing both scripts at the same
    env var silently never finds the model that was just trained."""
    monkeypatch.setattr(
        ecl, "fetch_groups", lambda user_id: [("BCP", "SOME NEW MERCHANT", 2)]
    )
    (tmp_path / "model.joblib").write_bytes(b"stand-in")
    seen_paths: list[Path] = []

    def _fake_load_model(path: Path) -> object:
        seen_paths.append(path)
        return object()

    monkeypatch.setattr(ecl, "load_model", _fake_load_model)
    monkeypatch.setattr(
        ecl, "suggest_from_model", lambda bundle, description, guess: "Deporte"
    )
    out = tmp_path / "labels.xlsx"

    assert (
        ecl.main(
            [
                "--user",
                "piero",
                "--out",
                str(out),
                "--model-path",
                str(tmp_path / "model"),  # no ".joblib"
            ]
        )
        == 0
    )

    assert seen_paths == [tmp_path / "model.joblib"]


def test_main_falls_back_to_rules_when_no_model_has_been_saved(
    environment: None, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    monkeypatch.setattr(
        ecl, "fetch_groups", lambda user_id: [("BCP", "NETFLIX.COM", 3)]
    )
    out = tmp_path / "labels.xlsx"

    assert (
        ecl.main(
            [
                "--user",
                "piero",
                "--out",
                str(out),
                "--model-path",
                str(tmp_path / "no-model-here.joblib"),
            ]
        )
        == 0
    )

    rows = list(load_workbook(out)["Categorias"].iter_rows(min_row=2, values_only=True))
    assert rows[0] == ("BCP", "NETFLIX.COM", 3, "Servicios", "Servicios")


def test_main_requires_a_user(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("PFP_USER", raising=False)

    assert ecl.main(["--out", "/tmp/x.xlsx"]) == 2


def test_main_refuses_to_write_an_empty_template(
    environment: None, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    monkeypatch.setattr(ecl, "fetch_groups", lambda user_id: [])

    assert ecl.main(["--user", "piero", "--out", str(tmp_path / "labels.xlsx")]) == 1
    assert not (tmp_path / "labels.xlsx").exists()
