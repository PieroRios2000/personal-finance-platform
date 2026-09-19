"""`pfp import-manual`: the manual Excel's savings into bronze, safe to run
again and again."""

from collections.abc import Sequence
from datetime import date
from pathlib import Path

import pytest
from deltalake import DeltaTable
from openpyxl import Workbook

from ingestion import cli
from ingestion.manual_excel import SAVINGS_COLUMNS

_ACCOUNT = "Ahorros Prueba"


@pytest.fixture(autouse=True)
def env(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    lake = tmp_path / "lake"
    monkeypatch.setenv("LAKEHOUSE_URI", str(lake))
    monkeypatch.setenv("PFP_ACCOUNT_KEY", "test-key")
    monkeypatch.setenv("PFP_USER", "piero")
    return lake


def _workbook(path: Path, rows: Sequence[tuple[object, ...]]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Ahorros"
    sheet.append(list(SAVINGS_COLUMNS))
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
    return path


_ROWS: list[tuple[object, ...]] = [
    (_ACCOUNT, date(2026, 7, 5), "deposito", 100.0, "PEN", 100.0),
    (_ACCOUNT, date(2026, 7, 20), "intereses", 1.5, "PEN", 101.5),
    (_ACCOUNT, date(2026, 8, 31), "cierre de mes", 0, "PEN", 101.5),
]


def _count(lake: Path, table: str) -> int:
    rows: int = DeltaTable(str(lake / "bronze" / table)).to_pyarrow_table().num_rows
    return rows


def test_the_savings_reach_bronze_and_the_report_has_counts_only(
    tmp_path: Path, env: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook = _workbook(tmp_path / "f.xlsx", _ROWS)

    code = cli.main(["import-manual", str(workbook)])

    out = capsys.readouterr().out
    assert code == 0
    assert "2 statement(s) written" in out
    assert _count(env, "statements") == 2
    assert _count(env, "transactions") == 2
    for private in (_ACCOUNT, "101.5", "100.0"):
        assert private not in out


def test_loading_the_same_workbook_again_adds_nothing(
    tmp_path: Path, env: Path
) -> None:
    workbook = _workbook(tmp_path / "f.xlsx", _ROWS)
    cli.main(["import-manual", str(workbook)])

    cli.main(["import-manual", str(workbook)])

    assert _count(env, "statements") == 2
    assert _count(env, "transactions") == 2


def test_a_corrected_month_replaces_the_old_one_and_new_months_are_added(
    tmp_path: Path, env: Path
) -> None:
    cli.main(["import-manual", str(_workbook(tmp_path / "a.xlsx", _ROWS))])
    corrected = [
        (_ACCOUNT, date(2026, 7, 5), "deposito", 100.0, "PEN", 100.0),
        (_ACCOUNT, date(2026, 7, 20), "intereses", 2.0, "PEN", 102.0),
        (_ACCOUNT, date(2026, 8, 31), "cierre de mes", 0, "PEN", 102.0),
        (_ACCOUNT, date(2026, 9, 10), "deposito", 10.0, "PEN", 112.0),
    ]

    cli.main(["import-manual", str(_workbook(tmp_path / "b.xlsx", corrected))])

    assert _count(env, "statements") == 3
    assert _count(env, "transactions") == 3
    amounts = sorted(
        str(a)
        for a in DeltaTable(str(env / "bronze" / "transactions"))
        .to_pyarrow_table()
        .column("amount")
        .to_pylist()
    )
    assert amounts == ["10.00", "100.00", "2.00"]


def test_nothing_is_written_when_the_workbook_has_problems(
    tmp_path: Path, env: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    bad = [*_ROWS, (_ACCOUNT, date(2026, 8, 31), "x", 1.0, "PEN", 999.0)]

    code = cli.main(["import-manual", str(_workbook(tmp_path / "f.xlsx", bad))])

    assert code == 1
    assert "balance does not follow" in capsys.readouterr().err
    assert not (env / "bronze").exists()


def test_a_missing_file_a_missing_user_and_a_missing_key_are_clear_errors(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    assert cli.main(["import-manual", str(tmp_path / "nope.xlsx")]) == 2
    assert "not found" in capsys.readouterr().err

    workbook = _workbook(tmp_path / "f.xlsx", _ROWS)
    monkeypatch.delenv("PFP_USER")
    assert cli.main(["import-manual", str(workbook)]) == 2
    assert "--user" in capsys.readouterr().err

    monkeypatch.setenv("PFP_USER", "piero")
    monkeypatch.delenv("PFP_ACCOUNT_KEY")
    assert cli.main(["import-manual", str(workbook)]) == 1
    assert "PFP_ACCOUNT_KEY" in capsys.readouterr().err


_INVEST_HEADER = ["lugar", "fecha", "tipo", "monto", "moneda", "saldo_final", "nota"]


def _with_investments(path: Path, rows: Sequence[tuple[object, ...]]) -> Path:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    sheet = workbook.create_sheet("Inversiones")
    sheet.append(_INVEST_HEADER)
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
    return path


_INVEST: list[tuple[object, ...]] = [
    ("Fondo A", date(2026, 7, 5), "aporte", 100.0, "PEN", 100.0, "Detalle Secreto"),
    ("Fondo A", date(2026, 7, 31), "valorizacion", 0, "PEN", 102.0, None),
    ("Fondo A", date(2026, 8, 31), "valorizacion", 0, "PEN", 103.0, None),
]


def test_the_investments_sheet_is_loaded_too_with_counts_only(
    tmp_path: Path, env: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook = _with_investments(_workbook(tmp_path / "f.xlsx", _ROWS), _INVEST)

    code = cli.main(["import-manual", str(workbook)])

    out = capsys.readouterr().out
    assert code == 0
    assert "Inversiones: 2 month(s) written" in out
    assert _count(env, "investment_entries") == 3
    for private in ("Fondo A", "Detalle Secreto", "102"):
        assert private not in out


def test_loading_the_investments_again_replaces_the_months_it_holds(
    tmp_path: Path, env: Path
) -> None:
    workbook = _with_investments(_workbook(tmp_path / "f.xlsx", _ROWS), _INVEST)
    cli.main(["import-manual", str(workbook)])
    cli.main(["import-manual", str(workbook)])

    assert _count(env, "investment_entries") == 3

    corrected = _with_investments(
        _workbook(tmp_path / "g.xlsx", _ROWS),
        [
            ("Fondo A", date(2026, 7, 5), "aporte", 100.0, "PEN", 100.0, None),
            ("Fondo A", date(2026, 7, 31), "valorizacion", 0, "PEN", 101.0, None),
            ("Fondo A", date(2026, 8, 31), "valorizacion", 0, "PEN", 103.0, None),
            ("Fondo A", date(2026, 9, 30), "valorizacion", 0, "PEN", 104.0, None),
        ],
    )
    cli.main(["import-manual", str(corrected)])

    assert _count(env, "investment_entries") == 4


def test_a_problem_in_the_investments_stops_the_savings_from_loading_too(
    tmp_path: Path, env: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    bad = [("Fondo A", date(2026, 7, 5), "compra", 100.0, "PEN", 100.0, None)]
    workbook = _with_investments(_workbook(tmp_path / "f.xlsx", _ROWS), bad)

    code = cli.main(["import-manual", str(workbook)])

    assert code == 1
    assert "Inversiones row 2: tipo must be" in capsys.readouterr().err
    assert not (env / "bronze").exists()


def test_months_without_a_valuation_and_missing_months_are_reported_as_counts(
    tmp_path: Path, env: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    rows: list[tuple[object, ...]] = [
        ("Fondo A", date(2026, 7, 5), "aporte", 100.0, "PEN", 100.0, None),
        ("Fondo A", date(2026, 9, 5), "aporte", 100.0, "PEN", 205.0, None),
    ]
    workbook = _with_investments(_workbook(tmp_path / "f.xlsx", _ROWS), rows)

    cli.main(["import-manual", str(workbook)])

    out = capsys.readouterr().out
    assert "Inversiones: 2 month(s) without a valuation" in out
    assert "Inversiones: 1 month(s) missing between the first and last" in out


def test_a_workbook_without_an_investments_sheet_says_nothing_was_loaded(
    tmp_path: Path, env: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    cli.main(["import-manual", str(_workbook(tmp_path / "f.xlsx", _ROWS))])

    assert "Inversiones: nothing to load" in capsys.readouterr().out


def test_an_unreadable_workbook_is_reported_once(
    tmp_path: Path, env: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    path = tmp_path / "f.xlsx"
    path.write_text("not a zip")

    code = cli.main(["import-manual", str(path)])

    assert code == 1
    assert capsys.readouterr().err.count("not a readable .xlsx workbook") == 1
