"""Describes the manual Excel without exposing what is in it (ADR 0004).

The workbook holds real balances and movements, so this prints only structure
and counts, the same rule as `inspect_pdf_layout.py`: text cells come out as
shapes (letters `X`, digits `9`), numbers and dates never appear at all (only
how many there are), and a group's name (an account or a fund) is never printed.
Column names are shown when they are the agreed ones and masked otherwise. Never
prints the file's path.

    uv run python -m scripts.inspect_manual_excel <workbook.xlsx>
"""

import argparse
import string
import sys
from collections import Counter
from collections.abc import Sequence
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ingestion.manual_excel import SAVINGS_COLUMNS
from scripts.make_manual_templates import INVESTMENT_COLUMNS

_EXPECTED = {"Ahorros": SAVINGS_COLUMNS, "Inversiones": INVESTMENT_COLUMNS}
_KNOWN_TYPES = ("aporte", "retiro", "valorizacion")
_MONTH_END = "cierre de mes"
_TOP_SHAPES = 8


def _shape(text: str) -> str:
    return "".join(
        "9" if c.isdigit() else c if c in string.punctuation or c == " " else "X"
        for c in text
    )


def _month(value: Any) -> tuple[int, int] | None:
    return (value.year, value.month) if isinstance(value, (date, datetime)) else None


def _cents(value: Any) -> Decimal | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return Decimal(str(round(float(value), 2)))


def _groups(rows: list[dict[str, Any]], key: str) -> dict[Any, list[dict[str, Any]]]:
    grouped: dict[Any, list[dict[str, Any]]] = {}
    for row in rows:
        grouped.setdefault(row.get(key), []).append(row)
    return grouped


def _column_lines(name: str, values: list[Any]) -> list[str]:
    types = Counter(type(v).__name__ for v in values if v is not None)
    empty = sum(1 for v in values if v is None)
    lines = [
        f"column {name}: types "
        + ", ".join(f"{t}={n}" for t, n in sorted(types.items()))
        + (f"; empty={empty}" if empty else "")
    ]
    texts = Counter(_shape(v) for v in values if isinstance(v, str))
    for shape, count in texts.most_common(_TOP_SHAPES):
        lines.append(f"  {shape} ({count})")
    numbers = [_cents(v) for v in values]
    numeric = [n for n in numbers if n is not None]
    if numeric:
        lines.append(
            f"  numbers: zero={sum(1 for n in numeric if n == 0)}, "
            f"negative={sum(1 for n in numeric if n < 0)}"
        )
    return lines


def _savings_lines(rows: list[dict[str, Any]], names: Sequence[str]) -> list[str]:
    lines: list[str] = []
    checked = followed = 0
    months: list[str] = []
    for group in _groups(rows, names[0]).values():
        months.append(str(len({m for r in group if (m := _month(r.get("fecha")))})))
        previous: Decimal | None = None
        for row in group:
            amount, balance = _cents(row.get("monto")), _cents(row.get("saldo_final"))
            if previous is not None and amount is not None and balance is not None:
                checked += 1
                followed += int(balance == previous + amount)
            previous = balance
    lines.append(
        f"balance chain: {followed} of {checked} rows follow the previous balance"
    )
    lines.append(f"months covered per group: {', '.join(months) or '0'}")
    markers = sum(
        1 for r in rows if str(r.get("descripcion") or "").strip().lower() == _MONTH_END
    )
    lines.append(f"cierre de mes rows: {markers}")
    return lines


def _investment_lines(rows: list[dict[str, Any]], names: Sequence[str]) -> list[str]:
    counts = Counter(str(r.get("tipo")) for r in rows)
    known = [f"{t}={counts[t]}" for t in _KNOWN_TYPES if counts[t]]
    other = sum(n for t, n in counts.items() if t not in _KNOWN_TYPES)
    if other:
        known.append(f"other={other}")
    valued: list[str] = []
    for group in _groups(rows, names[0]).values():
        any_row = {m for r in group if (m := _month(r.get("fecha")))}
        with_valuation = {
            m
            for r in group
            if r.get("tipo") == "valorizacion" and (m := _month(r.get("fecha")))
        }
        valued.append(f"{len(with_valuation)} of {len(any_row)}")
    return [
        f"tipo: {', '.join(known)}",
        f"months with a valuation per group: {', '.join(valued) or '0 of 0'}",
    ]


def _describe_sheet(sheet: Worksheet, index: int) -> list[str]:
    known = sheet.title in _EXPECTED
    title = sheet.title if known else _shape(sheet.title)
    expected = _EXPECTED.get(sheet.title, ())
    header = [c.value for c in sheet[1]]
    names = [
        h if isinstance(h, str) and h in expected else _shape(str(h))
        for h in header
        if h is not None
    ]
    body = list(sheet.iter_rows(min_row=2, values_only=True))
    body = [row for row in body if any(v is not None for v in row)]
    lines = [f"== sheet {index}: {title} ==", f"rows: {len(body)}"]
    lines.append("columns: " + " | ".join(names))
    records = [dict(zip(names, row, strict=False)) for row in body]
    for position, name in enumerate(names):
        lines += _column_lines(
            name, [row[position] for row in body if position < len(row)]
        )
    if known and tuple(names) == expected:
        extra = (
            _savings_lines(records, names)
            if sheet.title == "Ahorros"
            else _investment_lines(records, names)
        )
        lines += extra
    elif known:
        lines.append("(columns differ from the template: no derived checks)")
    return lines


def describe(path: Path) -> str:
    workbook = load_workbook(path, data_only=True)
    lines = [
        "sheets: "
        + ", ".join(s if s in _EXPECTED else _shape(s) for s in workbook.sheetnames)
    ]
    for index, sheet in enumerate(workbook.worksheets, start=1):
        if sheet.title == "Instrucciones":
            continue
        lines += _describe_sheet(sheet, index)
    return "\n".join(lines)


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="inspect_manual_excel")
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args(argv)
    if not args.workbook.exists():
        print("inspect_manual_excel: file not found", file=sys.stderr)
        return 2
    print(describe(args.workbook))
    return 0


if __name__ == "__main__":
    sys.exit(main())
