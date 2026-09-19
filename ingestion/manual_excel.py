"""Reads the manual Excel's `Ahorros` sheet (savings accounts, Banco Ripley
first) into the same `Statement`s the PDF parsers produce, so they flow through
bronze, silver and gold and are reconciled the same way.

Banco Ripley gives no statements, so the owner types the movements (ADR 0027,
docs/manual-data.md). Rules that follow from that:

- One statement per account, currency and calendar month. A month with no
  movements is a `cierre de mes` row (amount 0): a balance marker, not a
  transaction (`Transaction` rejects a zero amount, ADR 0005).
- Each row's `saldo_final` must be the previous row's plus or minus its
  `monto`. That is the reconciliation for this source: a mistyped balance or
  amount is reported, not absorbed. An amount typed as a positive number whose
  balance went down is a withdrawal (people type them unsigned); a negative
  amount is taken as typed. The first row of an account has no previous balance
  and is taken as typed. Between months, dbt's continuity test checks the same thing.
- A month's identity (the `file_sha256` bronze keys it on) comes from user,
  account, currency and month, **not** from its content, so loading a corrected
  workbook *replaces* that month (`bronze.replace_statement`) and loading the
  same one again changes nothing.
- Problems are reported as row numbers and column names, never as values.
- The account is identified by its name in the sheet (there is no account
  number to hash), and `account_last4` is `0000`.

`ingestion` never imports `lakehouse` (only `ingestion.cli` does): this module
returns statements and the CLI writes them.
"""

import calendar
import hashlib
import zipfile
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ingestion.reconciliation import reconcile
from ingestion.schema import (
    Currency,
    InvestmentEntry,
    InvestmentKind,
    InvestmentMonth,
    Statement,
    Transaction,
    hash_account,
    normalize_description,
)

SAVINGS_COLUMNS = ("cuenta", "fecha", "descripcion", "monto", "moneda", "saldo_final")
INVESTMENT_COLUMNS = (
    "lugar",
    "fecha",
    "tipo",
    "monto",
    "moneda",
    "saldo_final",
    "nota",
)
SHEET = "Ahorros"
INVESTMENT_SHEET = "Inversiones"
_EXAMPLE_PREFIX = "EJEMPLO"
_MONTH_END = "cierre de mes"


@dataclass(frozen=True)
class Entry:
    statement: Statement
    file_sha256: str


@dataclass
class SavingsImport:
    entries: list[Entry] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)
    missing_months: int = 0


@dataclass(frozen=True)
class _Row:
    number: int
    account: str
    day: date
    description: str
    amount: Decimal
    currency: Currency
    balance: Decimal


_FLOAT_NOISE = Decimal("0.000001")


def _money(value: Any) -> tuple[Decimal | None, str | None]:
    """The number as an amount with at most 2 decimals, or why it is not one.
    Excel stores floats, so binary noise (0.1 + 0.2) is rounded away, but a
    real third decimal (1.005) is reported, like everywhere else in the
    platform (`Transaction` rejects it too)."""
    if value is None:
        return None, (
            "is empty (a formula without a saved value? "
            "open and save the file in Excel)"
        )
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None, "is not a number"
    exact = Decimal(repr(float(value)))
    rounded = exact.quantize(Decimal("0.01"))
    if abs(exact - rounded) > _FLOAT_NOISE:
        return None, "has more than 2 decimals"
    return rounded, None


def _day(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    return value if isinstance(value, date) else None


def _parse_rows(
    raw: list[tuple[int, tuple[Any, ...]]],
) -> tuple[list[_Row], list[tuple[int, str]]]:
    rows: list[_Row] = []
    problems: list[tuple[int, str]] = []
    for number, cells in raw:
        account, when, description, amount, currency, balance = cells[:6]
        day = _day(when)
        value, value_reason = _money(amount)
        closing, closing_reason = _money(balance)
        found: list[str] = []
        if not (isinstance(account, str) and account.strip()):
            found.append("cuenta is empty")
        if day is None:
            found.append("fecha is not a date")
        if not (isinstance(description, str) and description.strip()):
            found.append("descripcion is empty")
        if value is None:
            found.append(f"monto {value_reason}")
        if currency not in ("PEN", "USD"):
            found.append("moneda must be PEN or USD")
        if closing is None:
            found.append(f"saldo_final {closing_reason}")
        if found:
            problems += [(number, message) for message in found]
            continue
        assert day is not None and value is not None and closing is not None
        rows.append(
            _Row(
                number,
                str(account).strip(),
                day,
                str(description).strip(),
                value,
                "PEN" if currency == "PEN" else "USD",
                closing,
            )
        )
    return rows, problems


def _identity(
    user_id: str, account_id: str, currency: str, year: int, month: int
) -> str:
    key = f"manual-excel|{user_id}|{account_id}|{currency}|{year}-{month:02d}"
    return hashlib.sha256(key.encode()).hexdigest()


def _months_between(first: tuple[int, int], last: tuple[int, int]) -> int:
    return (last[0] - first[0]) * 12 + last[1] - first[1] + 1


def read_savings(path: Path, *, user_id: str) -> SavingsImport:
    result = SavingsImport()
    try:
        workbook = load_workbook(path, data_only=True)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile):
        result.problems.append("the file is not a readable .xlsx workbook")
        return result
    if SHEET not in workbook.sheetnames:
        result.problems.append(f"the workbook has no '{SHEET}' sheet")
        return result
    sheet = workbook[SHEET]
    header = tuple(c.value for c in sheet[1])
    if header[: len(SAVINGS_COLUMNS)] != SAVINGS_COLUMNS:
        result.problems.append(f"{SHEET}: the columns are not the template's")
        return result

    raw = [
        (index, tuple(row))
        for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2)
        if any(v is not None for v in row)
    ]
    if not raw:
        result.problems.append(f"{SHEET}: the sheet has no rows")
        return result
    examples = sum(
        1
        for _, cells in raw
        if isinstance(cells[0], str) and cells[0].startswith(_EXAMPLE_PREFIX)
    )
    if examples:
        result.problems.append(
            f"{SHEET}: {examples} example row(s) are still in the sheet (EJEMPLO)"
        )
        return result

    rows, row_problems = _parse_rows(raw)
    found: list[tuple[int, str]] = list(row_problems)
    groups: dict[tuple[str, Currency], list[_Row]] = {}
    for row in rows:
        groups.setdefault((row.account, row.currency), []).append(row)

    entries: list[Entry] = []
    for (account, currency), group in groups.items():
        group = sorted(group, key=lambda r: r.day)  # stable: sheet order within a day
        signed, chain_problems = _follow_balances(group)
        found += chain_problems
        if chain_problems:
            continue
        account_id = hash_account(account, account)
        by_month: dict[tuple[int, int], list[_Row]] = {}
        for row in signed:
            by_month.setdefault((row.day.year, row.day.month), []).append(row)
        result.missing_months += _months_between(min(by_month), max(by_month)) - len(
            by_month
        )
        for (year, month), month_rows in sorted(by_month.items()):
            sha = _identity(user_id, account_id, currency, year, month)
            try:
                statement = _month_statement(
                    user_id, account, account_id, currency, year, month, month_rows, sha
                )
            except ValueError:  # pydantic and reconciliation errors quote values
                found.append(
                    (
                        month_rows[0].number,
                        "the month could not be built (invalid values)",
                    )
                )
                continue
            entries.append(Entry(statement, sha))

    found.sort()
    result.problems += [f"{SHEET} row {row}: {message}" for row, message in found]
    if not result.problems:
        result.entries = entries
    return result


def _follow_balances(group: list[_Row]) -> tuple[list[_Row], list[tuple[int, str]]]:
    """Check each row's balance against the previous one and fix the sign of an
    amount typed unsigned. Returns the rows with signed amounts and the problems."""
    previous: Decimal | None = None
    signed: list[_Row] = []
    problems: list[tuple[int, str]] = []
    for row in group:
        amount = row.amount
        if previous is not None:
            if row.balance == previous + amount:
                pass
            elif amount > 0 and row.balance == previous - amount:
                # Typed unsigned (a withdrawal as a positive number): the
                # balance says which way the money went.
                amount = -amount
            else:
                problems.append(
                    (row.number, "the balance does not follow the previous one")
                )
        signed.append(replace(row, amount=amount))
        previous = row.balance
    return signed, problems


def _month_statement(
    user_id: str,
    account: str,
    account_id: str,
    currency: Currency,
    year: int,
    month: int,
    month_rows: list[_Row],
    sha: str,
) -> Statement:
    first, last = month_rows[0], month_rows[-1]
    statement = Statement(
        user_id=user_id,
        bank=account,
        account_id=account_id,
        account_last4="0000",
        period_start=date(year, month, 1),
        period_end=date(year, month, calendar.monthrange(year, month)[1]),
        opening_balance=first.balance - first.amount,
        closing_balance=last.balance,
        account_kind="asset",
        currency=currency,
        transactions=[
            Transaction(
                user_id=user_id,
                bank=account,
                account_id=account_id,
                account_last4="0000",
                date=r.day,
                description=normalize_description(r.description),
                amount=r.amount,
                currency=currency,
                source_file_sha256=sha,
            )
            for r in month_rows
            if r.amount != 0
        ],
    )
    reconcile(statement)
    return statement


@dataclass
class InvestmentImport:
    months: list[InvestmentMonth] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)
    months_without_valuation: int = 0
    missing_months: int = 0


_KINDS: tuple[InvestmentKind, ...] = ("aporte", "retiro", "valorizacion")


@dataclass(frozen=True)
class _InvestmentRow:
    number: int
    place: str
    day: date
    kind: InvestmentKind
    amount: Decimal
    currency: Currency
    balance: Decimal
    detail: str | None


def _investment_identity(
    user_id: str, place: str, currency: str, year: int, month: int
) -> str:
    key = f"manual-excel-investments|{user_id}|{place}|{currency}|{year}-{month:02d}"
    return hashlib.sha256(key.encode()).hexdigest()


def _parse_investment_rows(
    raw: list[tuple[int, tuple[Any, ...]]],
) -> tuple[list[_InvestmentRow], list[tuple[int, str]]]:
    rows: list[_InvestmentRow] = []
    problems: list[tuple[int, str]] = []
    for number, cells in raw:
        place, when, kind, amount, currency, balance, note = cells[:7]
        day = _day(when)
        value, value_reason = _money(amount)
        closing, closing_reason = _money(balance)
        found: list[str] = []
        if not (isinstance(place, str) and place.strip()):
            found.append("lugar is empty")
        if day is None:
            found.append("fecha is not a date")
        if kind not in _KINDS:
            found.append("tipo must be aporte, retiro or valorizacion")
        if value is None:
            found.append(f"monto {value_reason}")
        elif kind == "valorizacion" and value != 0:
            found.append("monto must be 0 in a valorizacion")
        elif kind in ("aporte", "retiro") and value <= 0:
            found.append("monto must be greater than 0 in an aporte or retiro")
        if currency not in ("PEN", "USD"):
            found.append("moneda must be PEN or USD")
        if closing is None:
            found.append(f"saldo_final {closing_reason}")
        elif closing < 0:
            found.append("saldo_final cannot be negative")
        if found:
            problems += [(number, message) for message in found]
            continue
        assert day is not None and value is not None and closing is not None
        rows.append(
            _InvestmentRow(
                number,
                str(place).strip(),
                day,
                "aporte"
                if kind == "aporte"
                else "retiro"
                if kind == "retiro"
                else "valorizacion",
                value,
                "PEN" if currency == "PEN" else "USD",
                closing,
                str(note).strip() if isinstance(note, str) and note.strip() else None,
            )
        )
    return rows, problems


def read_investments(path: Path, *, user_id: str) -> InvestmentImport:
    """Read the `Inversiones` sheet into one `InvestmentMonth` per fund, currency
    and calendar month (ADR 0028). The sheet is optional: a workbook without it,
    or with only the header, imports nothing and is not a problem."""
    result = InvestmentImport()
    try:
        workbook = load_workbook(path, data_only=True)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile):
        result.problems.append("the file is not a readable .xlsx workbook")
        return result
    if INVESTMENT_SHEET not in workbook.sheetnames:
        return result
    sheet = workbook[INVESTMENT_SHEET]
    header = tuple(c.value for c in sheet[1])
    if header[: len(INVESTMENT_COLUMNS)] != INVESTMENT_COLUMNS:
        result.problems.append(
            f"{INVESTMENT_SHEET}: the columns are not the template's"
        )
        return result
    raw = [
        (index, tuple(row))
        for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2)
        if any(v is not None for v in row)
    ]
    examples = sum(
        1
        for _, cells in raw
        if len(cells) > 6
        and isinstance(cells[6], str)
        and cells[6].startswith(_EXAMPLE_PREFIX)
    )
    if examples:
        result.problems.append(
            f"{INVESTMENT_SHEET}: {examples} example row(s) are still in the sheet "
            "(EJEMPLO)"
        )
        return result

    rows, row_problems = _parse_investment_rows(raw)
    found: list[tuple[int, str]] = list(row_problems)
    groups: dict[tuple[str, Currency], list[_InvestmentRow]] = {}
    for row in rows:
        groups.setdefault((row.place, row.currency), []).append(row)

    months: list[InvestmentMonth] = []
    for (place, currency), group in groups.items():
        group = sorted(group, key=lambda r: (r.day, r.number))
        by_month: dict[tuple[int, int], list[_InvestmentRow]] = {}
        for row in group:
            by_month.setdefault((row.day.year, row.day.month), []).append(row)
        result.missing_months += _months_between(min(by_month), max(by_month)) - len(
            by_month
        )
        for (year, month), month_rows in sorted(by_month.items()):
            if not any(r.kind == "valorizacion" for r in month_rows):
                result.months_without_valuation += 1
            try:
                months.append(
                    InvestmentMonth(
                        user_id=user_id,
                        place=place,
                        currency=currency,
                        year=year,
                        month=month,
                        month_key=_investment_identity(
                            user_id, place, currency, year, month
                        ),
                        entries=[
                            InvestmentEntry(
                                date=r.day,
                                kind=r.kind,
                                amount=r.amount,
                                balance=r.balance,
                                detail=r.detail,
                                position=r.number,
                            )
                            for r in month_rows
                        ],
                    )
                )
            except ValueError:
                found.append(
                    (
                        month_rows[0].number,
                        "the month could not be built (invalid values)",
                    )
                )
    found.sort()
    result.problems += [
        f"{INVESTMENT_SHEET} row {row}: {message}" for row, message in found
    ]
    if not result.problems:
        result.months = months
    return result
