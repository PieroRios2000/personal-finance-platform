"""The plan workbook's layout and what a re-export keeps (T56, ADR 0048, spec 4.1).

Knows only the *shape* of the file, never a real description or amount: those are
read from gold and written by `scripts/export_plan.py`, on the owner's machine. Sheet
names are Spanish like the owner's other workbooks; headers are English like the
labeling file.
"""

from dataclasses import dataclass, replace
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from forecasting.fixed_expenses import WINDOW_MONTHS, Candidate

SHEET_INSTRUCTIONS = "Instrucciones"
SHEET_ITEMS = "Gastos fijos"
SHEET_META = "Meta"

KINDS = ("fixed", "variable", "ignore")
ITEM_COLUMNS = (
    "bank",
    "description",
    "currency",
    "category",
    "months_seen",
    "typical_amount",
    "proposed_kind",
    "kind",
    "expected_amount",
    "note",
)
META_DEFAULTS: dict[str, object] = {
    "goal_amount": None,
    "usd_to_pen": None,
    "emergency_months": 6,
    "emergency_basis": "all",
    "emergency_account": "Ripley",
    "target_date": None,
    "income_pen_override": None,
    "income_usd_override": None,
}
_META_HELP = {
    "goal_amount": "Tu meta de ahorro, en dólares (US$).",
    "usd_to_pen": "Soles que cuestan 1 dólar, p. ej. 3.75. Sirve para convertir.",
    "emergency_months": "Meses de gasto que quieres como fondo de emergencia.",
    "emergency_basis": "all = todos los gastos; fixed_only = solo los fijos.",
    "emergency_account": "Nombre exacto de la cuenta de tu fondo de emergencia.",
    "target_date": "Opcional: fecha en que quieres alcanzar la meta (AAAA-MM-DD).",
    "income_pen_override": "Opcional: ingreso mensual en soles (si no, el histórico).",
    "income_usd_override": "Opcional: ingreso mensual en dólares (si no, histórico).",
}
_INSTRUCTIONS = (
    "Plan de ahorro: qué completar",
    "",
    "1. Hoja «Gastos fijos»: el sistema propone qué gastos parecen fijos (casi el "
    "mismo monto casi todos los meses). Tú decides en la columna kind: fixed (fijo), "
    "variable (cambia cada mes) o ignore (excluir: una compra única o una "
    "transferencia a tus fondos de inversión).",
    "2. expected_amount: el monto mensual que esperas de ese gasto. Viene con el monto "
    "típico; cámbialo si sabes que va a subir o bajar.",
    "3. Las demás columnas (bank, description, currency, category, months_seen, "
    "typical_amount, proposed_kind, note) son del sistema: no las edites.",
    "4. Hoja «Meta»: goal_amount es tu meta en dólares; usd_to_pen son los soles "
    "que cuestan 1 dólar; emergency_months son los meses de gasto que quieres como "
    "fondo de emergencia (6 por defecto); emergency_basis: all cuenta todos los "
    "gastos, fixed_only solo los fijos; emergency_account es el nombre exacto de tu "
    "cuenta de emergencia (Ripley por defecto). Los campos opcionales pueden "
    "quedar vacíos.",
    "5. Guarda el archivo. Si vuelves a correr make export-plan se conservan tus "
    "elecciones y solo se agregan los gastos nuevos que aparezcan.",
    "6. El siguiente paso, make import-plan, llega con la tarea T57.",
)
_NOT_DETECTED = f"no longer detected in the last {WINDOW_MONTHS} months"


@dataclass(frozen=True)
class PlanItem:
    bank: str
    description: str
    currency: str
    category: str
    months_seen: int
    typical_amount: float
    proposed_kind: str
    kind: str
    expected_amount: float | None
    note: str = ""

    @property
    def key(self) -> tuple[str, str, str]:
        return (self.bank, self.description, self.currency)

    @property
    def owner_edited(self) -> bool:
        return (
            self.kind != self.proposed_kind
            or self.expected_amount != self.typical_amount
        )


@dataclass(frozen=True)
class PlanFile:
    items: tuple[PlanItem, ...]
    meta: dict[str, object]


def _proposal(candidate: Candidate) -> PlanItem:
    return PlanItem(
        bank=candidate.bank,
        description=candidate.description,
        currency=candidate.currency,
        category=candidate.category,
        months_seen=candidate.months_seen,
        typical_amount=candidate.typical_amount,
        proposed_kind=candidate.proposed_kind,
        kind=candidate.proposed_kind,
        expected_amount=candidate.typical_amount,
        note=candidate.note,
    )


def merge(candidates: list[Candidate], previous: PlanFile | None) -> PlanFile:
    """Fresh proposals for new and untouched rows; the owner's `kind` and
    `expected_amount` and every filled `Meta` value survive; a row that is no longer
    detected stays and says so; new candidates go after the existing rows."""
    proposals = {_proposal(c).key: _proposal(c) for c in candidates}
    if previous is None:
        return PlanFile(tuple(proposals.values()), dict(META_DEFAULTS))

    items: list[PlanItem] = []
    for old in previous.items:
        fresh = proposals.pop(old.key, None)
        if fresh is None:
            items.append(replace(old, note=_NOT_DETECTED))
        elif old.owner_edited:
            items.append(
                replace(fresh, kind=old.kind, expected_amount=old.expected_amount)
            )
        else:
            items.append(fresh)
    items.extend(proposals.values())

    meta = dict(META_DEFAULTS)
    meta.update(
        {k: v for k, v in previous.meta.items() if k in META_DEFAULTS and v is not None}
    )
    return PlanFile(tuple(items), meta)


def write_plan(path: Path, plan: PlanFile) -> None:
    workbook = Workbook()
    instructions = workbook.active
    assert instructions is not None
    instructions.title = SHEET_INSTRUCTIONS
    for line in _INSTRUCTIONS:
        instructions.append([line])
    instructions["A1"].font = Font(bold=True)
    instructions.column_dimensions["A"].width = 120

    items = workbook.create_sheet(SHEET_ITEMS)
    items.append(list(ITEM_COLUMNS))
    for item in plan.items:
        items.append([getattr(item, column) for column in ITEM_COLUMNS])
    for cell in items[1]:
        cell.font = Font(bold=True)
    items.freeze_panes = "A2"
    items.column_dimensions["B"].width = 50
    items.column_dimensions["J"].width = 40
    kinds = DataValidation(type="list", formula1='"fixed,variable,ignore"')
    items.add_data_validation(kinds)
    kinds.add(f"H2:H{max(2, len(plan.items) + 1)}")

    meta = workbook.create_sheet(SHEET_META)
    meta.append(["field", "value", "help"])
    for field in META_DEFAULTS:
        meta.append([field, plan.meta.get(field), _META_HELP[field]])
    for cell in meta[1]:
        cell.font = Font(bold=True)
    meta.column_dimensions["A"].width = 22
    meta.column_dimensions["B"].width = 16
    meta.column_dimensions["C"].width = 80
    basis = DataValidation(type="list", formula1='"all,fixed_only"')
    meta.add_data_validation(basis)
    basis.add(f"B{list(META_DEFAULTS).index('emergency_basis') + 2}")

    workbook.save(path)


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _number(value: object) -> float | None:
    return None if value is None or _text(value) == "" else float(str(value))


def read_plan(path: Path) -> PlanFile:
    workbook = load_workbook(path, data_only=True)
    for name in (SHEET_ITEMS, SHEET_META):
        if name not in workbook.sheetnames:
            raise ValueError(f"{path.name} has no sheet '{name}'")

    items: list[PlanItem] = []
    rows = workbook[SHEET_ITEMS].iter_rows(min_row=2, values_only=True)
    for row in rows:
        cells = dict(zip(ITEM_COLUMNS, row, strict=False))
        if not _text(cells["description"]):
            continue
        items.append(
            PlanItem(
                bank=_text(cells["bank"]),
                description=_text(cells["description"]),
                currency=_text(cells["currency"]),
                category=_text(cells["category"]),
                months_seen=int(float(str(cells["months_seen"] or 0))),
                typical_amount=float(str(cells["typical_amount"] or 0)),
                proposed_kind=_text(cells["proposed_kind"]),
                kind=_text(cells["kind"]),
                expected_amount=_number(cells["expected_amount"]),
                note=_text(cells["note"]),
            )
        )

    meta: dict[str, object] = dict.fromkeys(META_DEFAULTS)
    for field, value, *_ in workbook[SHEET_META].iter_rows(min_row=2, values_only=True):
        if field in meta and _text(value) != "":
            meta[str(field)] = value
    return PlanFile(tuple(items), meta)
